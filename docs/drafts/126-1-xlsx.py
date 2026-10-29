# GDD 126-1 표(마크다운)를 사용자용 엑셀로 만드는 스크립트 (마스터V). 사용: GDD 126-1 절을 .md 로 떼어 낸 뒤 python3 docs/drafts/126-1-xlsx.py 126-1.md 출력.xlsx (openpyxl). 표 순서 = 3단계·물리·특수·혼합·히든 5개
import re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

md = open(sys.argv[1], encoding='utf-8').read()
out = sys.argv[2]

def tables(text):
    """마크다운 표를 순서대로 [헤더, 행...] 목록으로."""
    res, cur = [], None
    for line in text.splitlines():
        if line.startswith('|'):
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            if all(re.fullmatch(r'-+', c) for c in cells):
                continue
            if cur is None:
                cur = []
                res.append(cur)
            cur.append([c.replace('**', '') for c in cells])
        else:
            cur = None
    return res

T = tables(md)
t3, phys, spec, mixed, hidden = T  # B · C물리 · C특수 · C혼합 · D

FONT = 'Malgun Gothic'
thin = Side(style='thin', color='BFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill('solid', fgColor='1F3864')
sec_fill = PatternFill('solid', fgColor='D9E1F2')
memo_fill = PatternFill('solid', fgColor='FFF2CC')
wrap = Alignment(wrap_text=True, vertical='top')

wb = Workbook()

def title(ws, text, sub):
    ws['A1'] = text
    ws['A1'].font = Font(name=FONT, size=14, bold=True)
    ws['A2'] = sub
    ws['A2'].font = Font(name=FONT, size=9, color='595959')

def put_table(ws, row, header, rows, widths, memo=True, section=None):
    if section:
        ws.cell(row=row, column=1, value=section).font = Font(name=FONT, size=11, bold=True)
        for c in range(1, len(header) + 2):
            ws.cell(row=row, column=c).fill = sec_fill
        row += 1
    hdr = header + (['사용자 메모(여기에 고칠 점)'] if memo else [])
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = border
    row += 1
    for r in rows:
        vals = r + ([''] if memo else [])
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=i, value=v)
            c.font = Font(name=FONT, size=10, bold=(i == 1 or i == 2 and header[0] == '#'))
            c.alignment = wrap
            c.border = border
            if memo and i == len(vals):
                c.fill = memo_fill
        row += 1
    for i, w in enumerate(widths + ([34] if memo else []), 1):
        cur = ws.column_dimensions[get_column_letter(i)].width
        ws.column_dimensions[get_column_letter(i)].width = max(cur or 0, w)
    return row + 1

SUB = 'GDD 126-1 초안 2판(블록 피해 = 특공 · 평타/블록 치명 반영, 디테일 확인 대기) · 수치 전부 임시값(provisional, 구현 때 측정해 맞춤) · 블록 이름 = 호로 / 사신(같으면 하나) · 노란 칸 = 사용자 메모'

# 1. 공통 틀
ws = wb.active
ws.title = '공통 틀'
title(ws, '진화체별 특성 칸 · 고유 블록 — 공통 틀 (2판)', SUB)
r = put_table(ws, 4, ['공격', '피해 계산', '치명'], [
    ['평타', '힘 ÷ 적 방어', '평타 치명(기본 5% · 상한 50%)'],
    ['블록 전부(흡수·기본·고유·콤보·공허·해방)', '지능 ÷ 적 정신력(특공)', '블록 치명(기본 5% · 상한 50%, 블록 보정 +4%p 포함)'],
    ['예외: 4단계 물리 노멀·혈·폭주형 고유 블록 3체인', '그 체인만 힘 ÷ 적 방어(물리)', '블록 치명'],
    ['적이 주는 피해', '그대로 — 적 공격 물리면 내 방어, 특수면 내 정신력으로 막음', '적 치명 그대로'],
], [40, 44, 44], section='피해 방식(사용자 확정) · 확률 더하기 표기 = +n%p')
r = put_table(ws, r, ['타입', '블록 성격'], [
    ['물리(조 / 권)', '블록 피해는 작고 평타를 강화한다'],
    ['혼합(교 / 참)', '평타와 블록 양쪽에 효과가 들어간다'],
    ['특수(탄 / 귀)', '평타는 장식, 블록 피해 위주로 싸운다'],
    ['계열 분류의 쓰임', '피해 계산에는 안 쓴다 — 진화를 어느 타입으로 미는가 + 고유 블록 성격만'],
], [26, 90], section='블록 성격(타입)')
r = put_table(ws, r, ['단계', '특성 칸', '들어가는 것', '고유 블록'], [
    ['1 호로 / 혼백', '0', '없음', '없음'],
    ['2 길리안 / 영술원 학생', '1', '성향 특성(진화 순간 노력치 1위)', '없음(콤보 블록이 그 몫)'],
    ['3 아쥬커스 / 평대원', '2', '성향 + 속성 특성(3단계 시트, 약하고 범용)', '진화체마다 1개'],
    ['4 바스트로데 / 수석 대원', '3', '성향 + 3단계 속성 특성(그대로 들고 감) + 진화체 특성', '진화체마다 1개'],
], [26, 9, 52, 24], section='단계별 특성 칸')
r = put_table(ws, r, ['성향 특성(노력치)', '효과'], [
    ['강인한 몸(체력)', '최대 체력 ↑'],
    ['괴력(힘)', '평타 피해 ↑ (바뀜: 옛 \'물리 공격 피해\')'],
    ['블록 가속(지능)', '블록이 더 빨리 생김 (블록 피해는 지능 능력치 자체가 정함)'],
    ['단단한 껍질(방어)', '받는 물리 피해 ↓'],
    ['흔들림 없는 마음(정신력)', '받는 특수 피해 ↓'],
    ['평타 가속(속도)', '평타가 빨라짐'],
], [26, 90], section='성향 특성 6종(이름 그대로, 설명만 맞춤)')
rules = [
    ['망한 진화체 · 실패형', '속성·진화체 특성 대신 성향 특성 2개(39장 3).'],
    ['아란칼화 / 참백도 구상화', '깨도 칸은 그대로. 레스렉시온 / 만해 고유 효과(83장 4)만 붙는다.'],
    ['고유 블록 나오는 법', '물공 블록 보장 생성원(49-11) 그대로 — 자기 타이머로 따로 나온다. 이 몸의 기본 블록은 막기·회복만.'],
    ['진화하면', '고유 블록은 새 진화체 것으로 바뀐다. 3단계 속성 특성은 4단계에도 남는다.'],
    ['체인 계열', '4단계 = 내 타입 계열(물리 = 조 / 권, 특수 = 탄 / 귀, 혼합 = 교 / 참). 3단계 = 3계열 고르게.'],
    ['물 = 회피(49-10)', '물 속성 특성 = 회피. 물 진화체의 물 블록 = 블록 치명 +10%p(이 몫만 상한 50% 넘음). 물 물리형은 평타 치명으로.'],
    ['3단계 속성 특성(49-10)', '4단계에서 어느 타입이 돼도 쓸모 있게 평타·블록 양쪽 또는 방어 쪽에만 건다. 세기는 39장의 약 60~70%.'],
]
put_table(ws, r, ['규칙', '내용'], rules, [26, 90], section='규칙')
ws.freeze_panes = 'A4'

# 2. 3단계
ws = wb.create_sheet('3단계 7종')
title(ws, '3단계 아쥬커스 / 평대원 — 7종', SUB)
put_table(ws, 4, t3[0], t3[1:], [30, 18, 46, 40])
ws.freeze_panes = 'A5'

# 3. 4단계 일반
ws = wb.create_sheet('4단계 일반 18종')
title(ws, '4단계 바스트로데 / 수석 대원 — 일반 18종', SUB)
r = put_table(ws, 4, phys[0], phys[1:], [10, 26, 44, 48], section='물리(조 / 권) 6 — 블록을 쓰면 평타가 강해짐 · 노멀~혈은 지금 5단계 구현분을 4단계로 옮김 · 사신 이름은 권(주먹)으로 바꾸는 제안')
r = put_table(ws, r, spec[0], spec[1:], [10, 26, 44, 48], section='특수(탄 / 귀) 6 — 블록끼리 연계, 특성은 단순')
r = put_table(ws, r, mixed[0], mixed[1:], [10, 26, 44, 48], section='혼합(교 / 참) 6 — 평타와 블록이 서로를 채움')
c = ws.cell(row=r, column=1, value='타입 제한(126장 6)으로 버리는 49-5 칸: 혈 마공·혈 혼합 · 독 물공·독 혼합 · 허 마공·허 혼합(허는 확정된 끌어당김을 혼합으로 옮김).')
c.font = Font(name=FONT, size=9, color='595959')

# 4. 4단계 히든
ws = wb.create_sheet('4단계 히든 10종')
title(ws, '4단계 히든 10종 — 조건은 GDD 126장 7·협의 답 그대로', SUB)
put_table(ws, 4, hidden[0], hidden[1:], [5, 16, 24, 50, 52])
ws.freeze_panes = 'A5'

# 5. 같이 바뀌는 것 · 정할 것
ws = wb.create_sheet('같이 바뀌는 것·정할 것')
title(ws, '같이 바뀌는 것 · 사용자가 정할 큰 방향', SUB)
r = put_table(ws, 4, ['항목', '내용'], [
    ['피해 공식', '평타 = 힘 ÷ 방어, 블록 = 지능 ÷ 정신력, 혼합 평균 삭제, 기본 공격 블록도 특공'],
    ['치명 두 값', '평타 치명 / 블록 치명, 상태 화면 두 줄'],
    ['특성 이름 바꿔 옮기기', '완성 → 4속성형 · 한계 돌파 → 폭주 본능 / 호로화 본능(폭주형) · 억눌린 힘 → 절제(절제형) · 사냥꾼 그대로 · 다시 일어남 → 각성형'],
    ['39장 속성 특성 표', '3단계 시트의 속성 특성(약하게)으로 대체, 3단계부터'],
    ['45-1 물 치명 상한 돌파', '물 블록 치명으로 옮김'],
    ['콤보 블록 4개 → 3개', '2단계 그랑 레이 세로 / 파도의 90 흑관 · 3단계 야수의 폭풍 / 월아 연참 · 4단계 세로 오스큐라스 포화 / 파도의 91 천번 흑관 · 맹수의 턱 / 철권 난타는 뺌'],
    ['희귀 특성', '삭제(128장 1)'],
    ['새 전투 효과', '회피·방벽·표식·물방울·작열·전기장·가열·무작위 피해 등(49-5 확정분 중 아직 안 만든 것 포함)'],
    ['밸런스', '1~3단계 가드(피해 공식 바뀐 몫 재측정), 4단계는 측정표만(45-6 후반 보류)'],
    ['세이브', '옛 5단계 → 같은 속성·타입 4단계 마이그레이션(MIN_LOADABLE 54 고정)'],
], [26, 100], memo=False, section='같이 바뀌는 것')
put_table(ws, r, ['#', '질문', '마스터 추천', '다른 안'], [
    ['1', '특성 칸', '0 · 1 · 2 · 3칸(단계마다 +1)', '4단계를 더 늘림'],
    ['2', '3단계 히든 통로 블록', '무작위 변이 1개(나올 때마다 물·불·흙·뇌 중 하나)', '히든 줄기마다 따로'],
    ['3', '사신 물리 블록 이름', '권(주먹) 이름으로 바꿈(기합 정권·물결 연권·잔불 주먹·전광 발차기·혈권)', '지금 칼 이름 그대로'],
], [5, 26, 60, 30], section='정할 큰 방향(답 없으면 추천대로, 답은 노란 칸에)')

for s in wb.worksheets:
    s.sheet_view.zoomScale = 100
wb.save(out)
print('saved', out, [s.title for s in wb.worksheets], len(t3)-1, len(phys)-1, len(spec)-1, len(mixed)-1, len(hidden)-1)
