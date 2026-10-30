# 게임 안에서 전용 아이콘/그림이 없는 것 목록 → docs/drafts/블리치_아이콘_없는_목록.xlsx
# 사용법: 저장소 루트에서  python3 docs/drafts/missing-icons-xlsx.py
# 데이터(json)·그림 폴더(public/art)·아이콘 스프라이트(src/ui/bha/sprite.svg)를 그때그때 읽는다.
# 화면 코드가 정하는 아이콘 규칙(PROP_ICON·상점/가방 아이콘·특성 아이콘)은 아래 표에 옮겨 적었다 — 코드가 바뀌면 같이 고칠 것.
#   src/ui/bha/room.ts PROP_ICON · src/ui/shopView.ts artOf · src/ui/lcdScreens.ts artOf · src/ui/buildPanel.ts traitIconId
#   src/ui/bha/pet.ts PET_ICON · src/ui/bha/expedition.ts tileIconId · src/ui/battleView.ts foeEl(foe-a/foe-b)
import json
import os
import re
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
OUT = os.path.join(ROOT, 'docs', 'drafts', '블리치_아이콘_없는_목록.xlsx')


def P(*a):
    return os.path.join(ROOT, *a)


def load(*a):
    with open(P(*a), encoding='utf-8') as f:
        return json.load(f)


ART = set(os.listdir(P('public', 'art')))
SPRITE = set(re.findall(r'<symbol id="([^"]+)"', open(P('src', 'ui', 'bha', 'sprite.svg'), encoding='utf-8').read()))


def has_file(key):
    stem = key.replace('.', '-')
    return f'{stem}.webp' in ART or f'{stem}.png' in ART or stem in ART  # 폴더(프레임 애니)도 그림 있음


def ic(i):
    """아이콘 이름 + 스프라이트에 없으면 표시"""
    return i if i in SPRITE else f'{i} (스프라이트에 없음!)'


FACTIONS = {f: load('src', 'data', 'factions', f'{f}.json') for f in ('hollow', 'shinigami')}
FNAME = {'hollow': '호로', 'shinigami': '사신', '*': '공통'}
ELEM = {e['id']: e['name']['*'] for e in load('src', 'data', 'elements', 'elements.json')}


def nm(d, f=None):
    """이름 사전 → 문자열. 종족별이면 '호로값 / 사신값'"""
    if isinstance(d, str):
        return d
    if f:
        return d.get(f) or d.get('*') or ''
    if '*' in d:
        return d['*']
    return ' / '.join(f'{FNAME[k]}: {v}' for k, v in d.items())


rows = []  # (분류, id, 이름, 지금 보이는 것, 종족, 비고)


def add(cat, id_, name, now, faction='공통', note=''):
    rows.append((cat, id_, name, now, faction, note))


# ───────────────────────── 1. 그림 키 (장소·물건·배경) ─────────────────────────
manifest_src = open(P('src', 'data', 'art', 'manifest.ts'), encoding='utf-8').read()
body = manifest_src.split('export const ART_MANIFEST', 1)[1].split('};', 1)[0]
mkeys = re.findall(r"'((?:scene|prop|place)\.[^']+)'", body)

# 화면 코드(room.ts PROP_ICON, faction.ts 기본 이름)에서 옮긴 값
PROP_ICON = {'altar': 'i-practice', 'book': 'i-shelf', 'chest': 'i-deck', 'market': 'i-flag-win', 'shop': 'i-heal', 'spa': 'i-tile-rest',
             'wanderer': 'i-special', 'gearShop': 'i-guard', 'skinShop': 'i-rare', 'traitShop': 'i-slot', 'smith': 'i-attack',
             'outside': 'i-loop', 'home': 'i-nest', 'adventure': 'i-loop', 'tower': 'i-title'}
PROP_NAME = {'altar': '수련장', 'book': '책장', 'chest': '가방(상자)', 'outside': '밖으로', 'door': '문(옛 모험 문)', 'market': '시장(월드맵 장소)',
             'shop': '소모품 상점', 'gearShop': '장비 상점', 'skinShop': '스킨 상점', 'traitShop': '특성 상점', 'smith': '대장장이',
             'wanderer': '떠돌이 상인', 'spa': '온천(장소)', 'home': '홈(둥지, 월드맵 장소)', 'adventure': '모험(월드맵 장소)'}
SCENE_NAME = {'cave': '호로 방(동굴) 배경', 'room': '사신 방 배경', 'spa': '온천 화면 배경', 'market': '시장(장터) 화면 배경', 'worldmap': '월드맵 하이앵글 배경'}
SCENE_NOW = {'market': '임시 배경(천막 줄·좌판 바닥 CSS 도형)', 'worldmap': '임시 지도 배경(CSS)'}
PROP_WHERE = {'altar': '방', 'book': '방', 'chest': '방', 'outside': '방', 'door': '방', 'market': '월드맵', 'home': '월드맵', 'adventure': '월드맵',
              'shop': '시장', 'gearShop': '시장', 'skinShop': '시장', 'traitShop': '시장', 'smith': '시장', 'wanderer': '시장', 'spa': '월드맵(온천)'}

for k in mkeys:
    if has_file(k):
        continue
    parts = k.split('.')
    if parts[0] == 'scene':
        add('배경(scene)', k, SCENE_NAME.get(parts[1], parts[1]), SCENE_NOW.get(parts[1], '임시 배경(CSS)'), '공통', '파일 이름 ' + k.replace('.', '-') + '.png')
    elif parts[0] == 'prop':
        p = parts[1]
        now = f'아이콘 {ic(PROP_ICON[p])}' if p in PROP_ICON else '자리표시 도형'
        note = f'{PROP_WHERE.get(p, "")} · 파일 {k.replace(".", "-")}.png'
        if p == 'spa':
            note += ' · 월드맵 온천은 place.spa.<종족> 도 대신 찾음'
        add('물건(prop)', k, PROP_NAME.get(p, p), now, '공통', note)
    else:  # place.<타일/장소>.<종족>
        p, f = parts[1], parts[2]
        if p.startswith('rest-'):
            tile = next(t for t in load('src', 'data', 'tiles', 'tiles.json') if t['id'] == p)
            add('장소(place)', k, f'휴식 칸 장소 — {nm(tile["name"], f)}', '임시 그림(빛나는 웅덩이 + i-tile-rest)', FNAME[f], '원정 걷기 앞쪽에 서는 장소 · 파일 ' + k.replace('.', '-') + '.png')
        else:
            label = {'tower': '탑(월드맵 장소)', 'spa': '온천(월드맵 장소)'}.get(p, p)
            if p == 'tower':
                label = f"탑 — {FACTIONS[f]['tower'].get('name', '탑') if isinstance(FACTIONS[f].get('tower'), dict) else '탑'}(월드맵 장소)"
            add('장소(place)', k, label, f'아이콘 {ic(PROP_ICON[p])}', FNAME[f], '월드맵 · prop.' + p + ' 이 있으면 그쪽 먼저 · 파일 ' + k.replace('.', '-') + '.png')

# 방 [밖으로]는 prop.outside 가 없을 때 prop.door 그림으로 대신 → 비고에 적기
for i, r in enumerate(rows):
    if r[1] == 'prop.outside' and has_file('prop.door'):
        rows[i] = (r[0], r[1], r[2], '문 그림(prop-door) 대신 씀', r[4], r[5] + ' · 전용 그림 아님')

# ───────────────────────── 2. 적 ─────────────────────────
for e in load('src', 'data', 'enemies', 'enemies.json'):
    if e.get('testOnly') or has_file(f'enemy.{e["id"]}'):
        continue
    cls = '강적' if e['class'] == 'elite' else '일반'
    add('적', e['id'], nm(e['name']), f'실루엣 {"foe-b" if e["class"] == "elite" else "foe-a"} + 속성 색 오라', '공통',
        f'{cls} · {ELEM.get(e["element"], e["element"])} · {e.get("category", "")} · 파일 enemy-{e["id"]}.png')
for b in load('src', 'data', 'bosses', 'bosses.json')['bosses']:
    if has_file(f'enemy.{b["id"]}'):
        continue
    add('적', b['id'], nm(b['name']), '실루엣 foe-b + 속성 색 오라', '종족별 이름', f'보스 · {ELEM.get(b["element"], b["element"])} · 파일 enemy-{b["id"]}.png')
for h in load('src', 'data', 'ranks', 'holders.json'):
    if has_file(f'enemy.{h["id"]}'):
        continue
    add('적', h['id'], nm(h['name']), f'실루엣 {"foe-b" if h.get("class") == "elite" else "foe-a"} + 속성 색 오라', '공통',
        f'등급전 보유자 {h["rank"]}위 · {ELEM.get(h["element"], h["element"])} · 파일 enemy-{h["id"]}.png')
for f, F in FACTIONS.items():
    add('적', f'ascension-foe ({f})', F['ascension']['foe'], '보스 forest-lord 그림 키를 같이 씀(지금은 실루엣 foe-b)', FNAME[f],
        f'{F["ascension"]["name"]} 전투 상대 — 전용 그림 키 없음(enemy-forest-lord 를 그리면 같이 바뀜)')

# ───────────────────────── 3. 캐릭터 진화체·스킨 ─────────────────────────
# 단계별 파일 char-<종족>-<n> 이 없을 때 실제로 보이는 것 (manifest.ts charArt 규칙)
for f, F in FACTIONS.items():
    tiers = list(F['tierNames'])
    for i, t in enumerate(tiers, start=1):
        key = f'char.{f}.{i}'
        if has_file(key):
            continue
        if f == 'hollow':
            if i == 1 and F.get('startLooks'):
                if all(has_file(f'skin.{s}') for s in F['startLooks']):
                    continue  # 시작 외형 스킨 4종이 덮는다
            if str(i) in F.get('tierLooks', {}):
                looks = set(F['tierLooks'][str(i)]['stat'].values()) | {p['look'] for p in F['tierLooks'][str(i)].get('pairs', [])}
                if all(has_file(f'skin.{s}') for s in looks):
                    continue  # 단계 외형 스킨이 덮는다
            now = '종족 공통 그림 char-hollow.png' if has_file('char.hollow') else '자리표시 도형'
            add('캐릭터 진화체', key, f'{FNAME[f]} {i}단계 {t}', now, FNAME[f], f'단계 전용 그림 없음 · 파일 char-{f}-{i}.png')
        else:
            rig = 'shinigami-student(학생복 종이인형)' if i <= 2 else 'shinigami-base(사신복 종이인형)'
            if i <= 2:
                continue  # 학생 단계는 학생복 종이인형이 전용
            add('캐릭터 진화체', key, f'{FNAME[f]} {i}단계 {t}', f'종족 공통 종이인형 {rig}', FNAME[f], f'단계 전용 그림 없음(3·4단계 같은 모습) · 파일 char-{f}-{i}.png 넣으면 그쪽이 먼저')
    asc = F['ascension']
    if asc.get('lookChange'):
        n = asc.get('artTier', len(tiers) + 1)
        add('캐릭터 진화체', f'char.{f}.{n}', f'{FNAME[f]} {asc["name"]} 후(아란칼)', '아란칼 기본형 종이인형 arrancar-base', FNAME[f],
            '진화체별 아란칼 그림은 아래 skin-arr* (아직 0장)')
    else:
        add('캐릭터 진화체', f'(외형 변화 없음) {f}', f'{FNAME[f]} {asc["name"]} 후', '그 전 단계 모습 그대로', FNAME[f], 'lookChange=false — 만해 모습 그림 자리 없음(참고)')

# 진화체(forms.json) — 진화체 하나하나의 전용 그림은 키 자체가 없다. 호로 4단계만 formLooks 로 스킨에 잇는다
forms = load('src', 'data', 'evolution', 'forms.json')
skin_names = {}
src_sk = open(P('src', 'data', 'skins', 'index.ts'), encoding='utf-8').read()
for m in re.finditer(r"\{ id: '([a-z]+)', name: '([^']+)', factionId: '(\w+)'", src_sk):
    skin_names[m.group(1)] = (m.group(2), m.group(3))
arr_el = re.findall(r"\['(\w+)', '([^']+)'\]", src_sk.split('ARRANCAR_ELEMENTS', 1)[1].split(';', 1)[0])
for eid, en in arr_el:
    for n, c in enumerate('abcd', start=1):
        skin_names[f'arr{eid}{c}'] = (f'아란칼 {en} {n}', 'hollow')
hid = ['폭주(frenzyName)', '절제', '강적 사냥꾼', '환생']
hid[0] = FACTIONS['hollow'].get('frenzyName', '폭주')
for i, c in enumerate('abcd'):
    skin_names[f'arrhidden{c}'] = (f'아란칼 {hid[i]}', 'hollow')
form_of_skin = defaultdict(list)
for fid, lst in FACTIONS['hollow'].get('formLooks', {}).items():
    for s in lst:
        form_of_skin[s].append(fid)

for sid, (sname, sf) in sorted(skin_names.items(), key=lambda x: (not x[0].startswith('arr'), x[0])):
    if has_file(f'skin.{sid}'):
        continue
    used = ', '.join(form_of_skin.get(sid, [])) or '진화체 표에 안 쓰임'
    add('스킨', f'skin.{sid}', sname, '아란칼 기본형 종이인형(arrancar-base)' if sid.startswith('arr') else '진화체 그림', FNAME.get(sf, sf),
        f'쓰는 진화체: {used} · 파일 skin-{sid}.png')

mapped = set(FACTIONS['hollow'].get('formLooks', {}).keys())
for fm in forms:
    if fm['tier'] == 4 and fm['id'] not in mapped:
        add('캐릭터 진화체', fm['id'], f'호로 4단계 진화체 {fm["label"]["kind"]} {"+".join(fm["label"].get("elements", []))}', '아란칼화 후 arrancar-base 종이인형(아란칼 외형 표 formLooks 에 없음)', '호로',
            '아란칼 외형 스킨이 배정되지 않은 진화체 — 그림 받으면 formLooks 에 이어야 함')
add('캐릭터 진화체', 'forms.json (사신 전부)', f'사신 진화체 {sum(1 for x in forms)}종', '단계 공통 종이인형', '사신', '사신은 진화체별 외형 표(formLooks·tierLooks)가 비어 있음 — 진화체마다 같은 모습')

# ───────────────────────── 4. 아이템류 ─────────────────────────
shop = load('src', 'data', 'shop', 'items.json')
MED = {'collapsed': 'i-heal', 'weak': 'i-hunger', 'tired': 'i-cond'}
bal = open(P('src', 'config', 'balance.ts'), encoding='utf-8').read()
heal_ids = set(re.findall(r"'([\w-]+)': \{ effect: 'heal'", bal))


def shop_icon(g):
    k = g['kind']
    if k == 'food': return 'i-hunger'
    if k == 'calm': return 'i-risk'
    if k == 'wake': return 'i-cond'
    if k == 'medicine': return MED[g['cures']]
    if g['id'] in heal_ids: return 'i-heal'
    if k in ('tileDraw', 'rareTile'): return 'i-tile-combo'
    if k == 'deckSlot': return 'i-deck'
    if k == 'reroll': return 'i-swap'
    return 'i-practice'


def bag_icon(g):
    if g.get('cures'): return MED[g['cures']]
    c, i = g['category'], g['id']
    if c == 'food': return 'i-hunger'
    if c == 'consumable':
        return 'i-risk' if ('calm' in i or i == 'elixir') else 'i-swap' if i == 'reroll' else 'i-cond'
    return 'i-practice'


GEAR_ICON = {k['id']: k['icon'] for k in load('src', 'data', 'gear', 'gear.json')['kinds']}
for g in shop['gold']:
    if g['kind'] == 'gear':
        continue
    if g['kind'] == 'food':
        name = ' / '.join(f'{FNAME[f]}: {F["foodNames"][g["food"]]}' for f, F in FACTIONS.items())
        cat = '음식'
    else:
        name, cat = g.get('name', g['id']), ('소모품(회복)' if g['id'] in heal_ids or g['kind'] == 'medicine' else '소모품·기타')
    s, b = shop_icon(g), bag_icon(g) if g['category'] in ('food', 'consumable', 'etc') else None
    now = f'상점 {ic(s)}' + (f' · 가방 {ic(b)}' if b and b != s else '')
    add(cat, g['id'], name, now, '종족별 이름' if g['kind'] == 'food' else '공통', f'{g["kind"]} · 전용 아이콘 필드 없음(종류별 공용 아이콘)')

gold_ids = {g['id'] for g in shop['gold']}
for w in load('src', 'data', 'shop', 'wanderer.json')['items']:
    if w['kind'] == 'gear':
        now = 'i-unknown + 등급 색'
    elif w['id'] in gold_ids:
        continue  # 위 상점 품목과 같은 것
    else:
        now = '상점 i-rare(떠돌이 상인 품목 공통)'
    add('떠돌이 상인 품목', w['id'], w['name'], now, '공통', f'{w["grade"]} · {w["kind"]}')

for f, F in FACTIONS.items():
    pass
for k in load('src', 'data', 'gear', 'gear.json')['kinds']:
    names = ' / '.join(f'{FNAME[f]}: {", ".join(sorted({v for s, v in F["gearSlotNames"].items() if (s if s in GEAR_ICON else "accessory") == k["id"]}))}' for f, F in FACTIONS.items())
    add('장비', k['id'], names, f'{ic(k["icon"])} + 등급 색', '종족별 이름', '장비는 개별 그림 없이 종류 아이콘 4개뿐(모든 장비·등급 공유)')

for m in load('src', 'data', 'materials', 'materials.json')['materials']:
    add('재료', m['id'], m['name'], ic(m['icon']), '공통', '빌려 쓰는 공용 UI 아이콘(재료 전용 아님)')

PET_ICON = {'gold': 'i-rare', 'belly': 'i-hunger', 'guard': 'i-guard', 'seek': 'i-target'}
for p in load('src', 'data', 'pets', 'pets.json')['pets']:
    if has_file(p['art']):
        continue
    name = ' / '.join(f'{FNAME[f]}: {F["petNames"][p["id"]]}' for f, F in FACTIONS.items())
    add('펫', p['art'], name, f'원 안에 효과 아이콘 {ic(PET_ICON[p["effect"]])}', '종족별 이름', f'파일 {p["art"].replace(".", "-")}.png')

for c in shop['cosmetics']:
    add('꾸미기(스킨 상점)', c['id'], c['name'], '색 견본(swatch, 토큰 색 칠한 도형)', '공통', f'{c["category"]} · 아이콘·그림 없음')
for d in shop['devices']:
    if has_file(f'device.{d["look"]}.front'):
        continue
    add('꾸미기(스킨 상점)', d['id'], d['name'], 'CSS 기기(그림 없음)', '공통', '시험용' if d.get('test') else '')

# ───────────────────────── 특성 ─────────────────────────
for t in load('src', 'data', 'traits', 'traits.json'):
    icons = []
    for f, F in FACTIONS.items():
        if t['kind'] != 'skill':
            i = 'i-rare' if t['kind'] == 'part' else 'i-slot'
        elif t.get('role') == 'guard':
            i = 'i-guard'
        elif t.get('role') == 'heal':
            i = 'i-heal'
        else:
            fx = next((x.get('fx') for x in F['families'] if x['id'] == t.get('family')), None)
            i = f'i-fx-{fx}' if fx else 'i-attack'
        icons.append(f'{FNAME[f]} {ic(i)}')
    add('흡수 특성', t['id'], nm(t['name']), ' / '.join(icons), '종족별 이름', f'{t["kind"]} · {ELEM.get(t.get("element"), t.get("element", ""))} · 계열 아이콘 공유(전용 없음)')
for t in load('src', 'data', 'traitShop', 'traits.json')['traits']:
    add('탈착 특성(특성 상점)', t['id'], t['name'], ic('i-slot'), '공통', '4개 모두 같은 아이콘')
ab = load('src', 'data', 'evolution', 'abilities.json')
fr = ' / '.join(F.get('frenzyName', '') for F in FACTIONS.values())
for grp, label in (('abilities', '기본'), ('elementAbilities', '속성'), ('uniqueAbilities', '고유')):
    for k, v in ab[grp].items():
        n = nm(v['name']).replace('{frenzyName}', fr)
        if 'name4' in v:
            n += f' (4단계: {nm(v["name4"])})'
        add('진화 특성', f'{grp}.{k}', n, '아이콘 없음(글자만) · 받을 때 결과 줄 i-rare', '공통', f'{label} 특성 · 아이콘 필드 없음')

# ───────────────────────── 타일 ─────────────────────────
TILE_ICON = {'prey-nest': 'i-tile-prey', 'elite-den': 'i-tile-elite', 'environment': 'i-tile-env', 'combo': 'i-tile-env', 'rest': 'i-tile-rest', 'event': 'i-tile-combo', 'special': 'i-tile-combo'}
for t in load('src', 'data', 'tiles', 'tiles.json'):
    fs = ', '.join(FNAME[k] for k in t['name'])
    add('지형 타일', t['id'], nm(t['name']), ic(TILE_ICON[t['kind']]), fs, f'{t["kind"]} · 타일 종류 아이콘 공유' + (' · 이벤트 전용 아이콘 없어 조합 아이콘 씀' if t['kind'] in ('event', 'special') else ''))

# ───────────────────────── 5. 스킬 블록 (64×64 도트 53장 예정) ─────────────────────────
ELS = ['normal', 'fire', 'water', 'earth', 'thunder', 'blood', 'poison', 'kyo']
FAM_KO = {'claw': '조', 'bite': '교', 'shot': '탄', 'slash': '참', 'impact': '권', 'beam': '귀'}
block_dir_files = set()
for d in (P('public', 'art'), P('public', 'art', 'blocks'), P('docs', 'design', 'user-art', 'block-guide')):
    if os.path.isdir(d):
        block_dir_files |= set(os.listdir(d))


def block_has(stem):
    return f'{stem}.png' in block_dir_files or f'{stem}.webp' in block_dir_files


for f, F in FACTIONS.items():
    for fam in F['families']:
        for e in ELS:
            stem = f'{f}-{fam["fx"]}-{e}'
            if block_has(stem):
                continue
            add('스킬 블록', stem, f'{FNAME[f]} {FAM_KO[fam["fx"]]}({fam["fx"]}) · {ELEM[e]}', f'계열 아이콘 {ic("i-fx-" + fam["fx"])} + 속성 색', FNAME[f], f'공격 블록 · 파일 {stem}.png (가이드 docs/design/user-art/block-guide)')
for r, rk, i in (('guard', '막기', 'i-guard'), ('heal', '회복', 'i-heal'), ('void', '공허', 'i-void'), ('combo', '콤보', 'i-chain-2'), ('release', '해방', 'i-special')):
    if block_has(f'role-{r}'):
        continue
    add('스킬 블록', f'role-{r}', f'역할 블록 · {rk}', f'역할 아이콘 {ic(i)}', '공통', f'파일 role-{r}.png (종족별로 하면 +5)')

# ───────────────────────── 6. 기타 (주석 '그림 대기'·자리표시) ─────────────────────────
add('기타', 'scene.spa 모션', '온천 들어가 앉는 모션', 'CSS 임시 연출(몸 아래를 물이 가림) + 돌계단 임시 도형', '공통', 'src/ui/spaView.ts 주석 "온천 모션은 사용자 그림 대기"')
for t in load('src', 'data', 'transcendents', 'transcendents.json'):
    add('기타', f'transcendent.{t["id"]}', f'초월체 {t["name"]}', '그림 없음(글자만)', '공통', 'src/data/transcendents/index.ts 주석 "외형은 그림 대기" · 그림 키 규칙도 아직 없음')
for f, F in FACTIONS.items():
    add('기타', f'emblem ({f})', f'{F["name"]} 종족 표지', f'이모지 {F["emblem"]}', FNAME[f], '시작 화면 종족 카드 — 로고 그림(logo-*)이 없을 때만 보임' if has_file(F['logo'].replace('-', '.', 1)) else '시작 화면 종족 카드')
add('기타', 'hero 자리표시', '전투·방 캐릭터 자리표시 도형', '둥근 몸 + 눈 둘(CSS)', '공통', '캐릭터 그림이 하나도 없을 때만 — 지금은 모든 단계에 그림/종이인형이 있어 실제로는 안 보임')

# ───────────────────────── 엑셀 쓰기 ─────────────────────────
HEAD = ['분류', 'id(키)', '화면 이름(한국어)', '지금 보이는 것', '종족', '비고']
SHEETS = [
    ('장소·물건·배경', ('배경(scene)', '물건(prop)', '장소(place)')),
    ('적', ('적',)),
    ('캐릭터·스킨', ('캐릭터 진화체', '스킨')),
    ('아이템', ('음식', '소모품(회복)', '소모품·기타', '떠돌이 상인 품목', '장비', '재료', '펫', '꾸미기(스킨 상점)')),
    ('특성·타일', ('흡수 특성', '탈착 특성(특성 상점)', '진화 특성', '지형 타일')),
    ('스킬 블록', ('스킬 블록',)),
    ('기타', ('기타',)),
]
bold = Font(bold=True)
fill = PatternFill('solid', fgColor='DDDDDD')
wb = Workbook()
wb.remove(wb.active)


def write(ws, head, data):
    ws.append(head)
    for c in ws[1]:
        c.font, c.fill = bold, fill
    for r in data:
        ws.append(list(r))
    ws.freeze_panes = 'A2'
    for i, h in enumerate(head, start=1):
        w = max([len(str(h))] + [len(str(r[i - 1])) for r in data]) if data else len(h)
        ws.column_dimensions[get_column_letter(i)].width = min(70, max(8, int(w * 1.6) + 2))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical='top', wrap_text=True)


for title, cats in SHEETS:
    data = [r for r in rows if r[0] in cats]
    write(wb.create_sheet(title), HEAD, data)

# 요약: 분류별 개수 + 공용 아이콘 공유 묶음
cnt = Counter(r[0] for r in rows)
summary = [(c, n) for _, cats in SHEETS for c in cats if (n := cnt.get(c, 0))]
ws = wb.create_sheet('요약')
write(ws, ['분류', '개수'], summary + [('합계', sum(cnt.values()))])
ws.append([])
ws.append(['많이 공유하는 아이콘(아이템·재료·특성·타일 등 "지금 보이는 것"에 나온 횟수)', '횟수'])
for c in ws[ws.max_row]:
    c.font = bold
share = Counter()
for r in rows:
    for i in set(re.findall(r'\bi-[a-z0-9-]+', r[3])):
        share[i] += 1
for i, n in share.most_common(15):
    ws.append([i, n])
ws.column_dimensions['A'].width = 60

wb.save(OUT)
print(OUT)
for c, n in summary:
    print(f'{c}\t{n}')
print('합계', sum(cnt.values()))
print('공유 상위', share.most_common(10))
