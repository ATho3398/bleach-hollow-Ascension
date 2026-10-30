# 사용자 대기 3건 엑셀 초안 만들기 (GDD 13장 '사용자가 알려줄 일', 마스터Z 2026-10-30)
# 실행: python3 docs/drafts/user-drafts-xlsx.py  → docs/drafts/*.xlsx
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
F = 'Arial'
YEL = PatternFill('solid', start_color='FFFF00')
HEAD = PatternFill('solid', start_color='D9D9D9')
EX = PatternFill('solid', start_color='F2F2F2')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def sheet(ws, title_lines, headers, rows, user_cols, widths, example=None):
    r = 1
    for t in title_lines:
        ws.cell(r, 1, t).font = Font(name=F, bold=(r == 1), size=12 if r == 1 else 10)
        r += 1
    ws.cell(r, 1, '노란 칸 = 사용자가 채우는 칸 · 회색 줄 = 예시(형식 참고용, 지워도 됨) · 나머지는 지금 게임 상태(참고)').font = Font(name=F, italic=True, size=9)
    r += 2
    hr = r
    for c, h in enumerate(headers, 1):
        cell = ws.cell(hr, c, h)
        cell.font = Font(name=F, bold=True); cell.fill = YEL if c in user_cols else HEAD
        cell.alignment = Alignment(wrap_text=True, vertical='center'); cell.border = BOX
    r = hr + 1
    if example:
        for c, v in enumerate(example, 1):
            cell = ws.cell(r, c, v); cell.font = Font(name=F, italic=True, color='808080'); cell.fill = EX; cell.border = BOX
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        r += 1
    for row in rows:
        for c in range(1, len(headers) + 1):
            v = row[c - 1] if c - 1 < len(row) else ''
            cell = ws.cell(r, c, v); cell.font = Font(name=F); cell.border = BOX
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if c in user_cols: cell.fill = YEL
        r += 1
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = ws.cell(hr + 1, 1)

# ---------- 1. 진화체 · 특성 초안 ----------
EL = {'normal': '노멀', 'water': '물', 'fire': '불', 'earth': '흙', 'thunder': '뇌', 'blood': '혈', 'poison': '독', 'kyo': '허'}
CH = {'physical': '평타캐', 'special': '블록캐', 'mixed': '하이브리드'}
HID4 = {'fusion': '이중 속성 히든', 'complete': '4속성형', 'fail': '실패형', 'awakened': '각성형', 'frenzy': '폭주형', 'restraint': '절제형', 'stable': '안정형', 'hunter': '사냥꾼', 'void': '허무형'}
forms = json.load(open(os.path.join(ROOT, 'src/data/evolution/forms.json'), encoding='utf-8'))
rows = []
for f in forms:
    if f['tier'] < 3: continue
    lb = f['label']; k = lb['kind']; els = '+'.join(EL.get(e, e) for e in lb.get('elements', [])) or '-'
    ch = CH.get(lb.get('category', ''), '-')
    if f['tier'] == 3:
        kind = {'element': '일반(속성)', 'fail': '망한 진화체', 'passage': '3단계 히든(통로)'}[k]
        now = {'fail': '허우적 — 노력치 얻는 양 ×1/4 (진화해도 유지, 각성형이 되면 사라짐)', 'passage': '변이의 기운 — 주는 피해 +5% · 받는 피해 −5%'}.get(k, '지금 연결된 속성 특성(임시)')
        route = '-'
    else:
        kind = '일반' if k == 'arrancar' else '히든 · ' + HID4.get(k, k)
        now = {'awakened': '다시 일어남 — 전투마다 한 번, 쓰러질 때 체력 30% 로 버팀 (망체 특성 대신)', 'fail': '없음(망체 특성 허우적이 남음)'}.get(k, '지금 연결된 대표 특성(임시)')
        route = (EL.get(lb['elements'][0], '') + ' 3단계 특성 → 이 특성으로 바뀜') if k == 'arrancar' and lb.get('elements') else ('3단계 히든 특성 → 이 특성으로 바뀜' if k not in ('fail', 'awakened') else '-')
    rows.append([f['id'], f['tier'], kind, els, ch, now, route, '', '', '', '', ''])
wb = Workbook(); ws = wb.active; ws.title = '진화체·특성'
sheet(ws, ['블리치 — 진화체 · 특성 초안 (사용자 작성용)',
           '규칙(GDD 147~151장): 특성은 3·4차 진화 때만 1개 · 4차 진화 때 3차 특성이 사라지고 4차 특성으로 바뀜(같은 루트면 윗단계, 예 물보라 → 폭포) · 받으면 영구 · 망체 특성은 유지 · 칸 = 기본 1 + 4단계 +1 + 환생당 +1(최대 4)',
           '캐릭터 성격: 평타캐(물리형) · 블록캐(특수형) · 하이브리드(혼합형) — 대화용 이름(GDD 152장)'],
      ['진화체 id', '단계', '종류', '속성', '캐릭터 성격', '지금 특성(임시)', '4차 때 바뀌는 길', '특성 이름', '특성 효과', '고유 블록 이름', '고유 블록 효과', '메모'],
      rows, {8, 9, 10, 11, 12}, [20, 6, 18, 10, 12, 36, 28, 16, 34, 16, 34, 24],
      example=['(예) t3:water', 3, '일반(속성)', '물', '-', '', '-', '물보라', '평타 치명 +8%p', '물결 가르기', '8초 평타 간격 −30%', '4차 물 → 폭포'])
wb.save(os.path.join(HERE, '블리치_진화체_특성_초안.xlsx'))

# ---------- 2. 대장장이 역할 ----------
wb = Workbook(); ws = wb.active; ws.title = '대장장이'
sheet(ws, ['블리치 — 대장장이 역할 (사용자 작성용)',
           '정해진 것: 대장장이 = 강화 · 제작 · 합성 · 분해 (사용자 2026-10-30). 장터의 대장장이 칸은 지금 "아직 문을 열지 않았다"만 나옴(GDD 153장).'],
      ['기능', '무엇을(대상)', '넣는 것(재료·골드)', '나오는 것(결과)', '성공 확률·실패하면', '제한(하루 횟수·단계 등)', '화면에서 어떻게', '메모'],
      [['강화'], ['제작'], ['합성'], ['분해']], {2, 3, 4, 5, 6, 7, 8}, [10, 22, 26, 24, 22, 22, 24, 24],
      example=['(예) 강화', '장비 1개', '강철 2 · 연마제 1 · 300골드', '장비 +1 (능력치 +5%)', '+5부터 실패 시 단계 그대로', '장비당 최대 +10', '장비 고르고 [강화]', ''])
wb.save(os.path.join(HERE, '블리치_대장장이_역할.xlsx'))

# ---------- 3. 기타 아이템(재료) 용도 ----------
MATS = [('두꺼운 가죽', '가죽'), ('유연한 가죽', '가죽'), ('가벼운 가죽', '가죽'), ('질긴 가죽', '가죽'), ('흑요석', '광물·보석'),
        ('가벼운 금속', '금속'), ('강철', '금속'), ('유연한 금속', '금속'), ('무른 금속', '금속'), ('호로 이빨', '몸 부위'),
        ('가면 조각', '몸 부위'), ('온전한 가면', '몸 부위'), ('금 조각', '금속'), ('은 조각', '금속'), ('염료', '보조'),
        ('반짝이는 보석', '광물·보석'), ('연마제', '보조'), ('윤활제', '보조'), ('질긴 실', '천·실'), ('천 조각', '천·실'),
        ('아교', '보조'), ('호박(보석)', '광물·보석')]
wb = Workbook(); ws = wb.active; ws.title = '기타 아이템'
sheet(ws, ['블리치 — 기타 아이템(재료 22종) 용도 (사용자 작성용)',
           '지금 상태: 가방 [기타] 창 + 22종 데이터만 있음, 얻는 곳·쓰는 곳 없음(GDD 154장). 분류 칸은 마스터 추정 — 틀리면 고쳐 주세요.'],
      ['재료', '분류(추정)', '얻는 곳', '쓰는 곳(강화·제작·합성·분해)', '희귀도', '상점 판매가(골드)', '메모'],
      [[n, c] for n, c in MATS], {2, 3, 4, 5, 6, 7}, [16, 12, 26, 30, 10, 14, 24],
      example=['(예) 강철', '금속', '★2 이상 원정 전투 보상(10%)', '강화 · 제작(무기)', '보통', '40', ''])
wb.save(os.path.join(HERE, '블리치_기타아이템_용도.xlsx'))
print('ok')
